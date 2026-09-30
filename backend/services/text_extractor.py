"""
Step 1 of the pipeline: turn an uploaded PDF or DOCX file into plain text.

An ATS can only score what it can read, so if text extraction fails the
resume would also fail in a real ATS.
"""

import io
import re
from pathlib import Path

import pdfplumber
from docx import Document

from backend.config import ALLOWED_EXTENSIONS, MAX_FILE_SIZE_BYTES, MAX_FILE_SIZE_MB


class FileError(Exception):
    """Raised when the uploaded file is invalid or unreadable.
    The message is safe to show directly to the user."""


def validate_file(file_bytes: bytes, filename: str) -> str:
    """Check size, extension and file signature. Returns the extension (".pdf"/".docx")."""
    if not file_bytes:
        raise FileError("The uploaded file is empty.")

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise FileError(f"File is larger than {MAX_FILE_SIZE_MB} MB.")

    extension = Path(filename or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise FileError("Unsupported file type. Please upload a PDF or DOCX file.")

    # Check the "magic bytes" at the start of the file so a renamed
    # .txt/.jpg can't pretend to be a PDF. (DOCX files are ZIP archives -> "PK".)
    if extension == ".pdf" and not file_bytes.startswith(b"%PDF"):
        raise FileError("This file does not look like a valid PDF.")
    if extension == ".docx" and not file_bytes.startswith(b"PK"):
        raise FileError("This file does not look like a valid DOCX.")

    return extension


def _extract_pdf(file_bytes: bytes) -> str:
    pages = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
            # Hyperlinks (e.g. a "GitHub" link) are stored separately from the text.
            for link in page.hyperlinks or []:
                uri = link.get("uri")
                if uri:
                    pages.append(uri)
    return "\n".join(pages)


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    lines = [p.text for p in doc.paragraphs]

    # Many resume templates put content inside tables.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.append(cell.text)

    # Hyperlink targets (LinkedIn/GitHub URLs behind link text).
    for rel in doc.part.rels.values():
        if "hyperlink" in rel.reltype and rel.is_external:
            lines.append(rel.target_ref)

    return "\n".join(lines)


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Validate the file and return its text content."""
    extension = validate_file(file_bytes, filename)

    try:
        text = _extract_pdf(file_bytes) if extension == ".pdf" else _extract_docx(file_bytes)
    except Exception as exc:  # corrupted / password-protected files
        raise FileError(f"Could not read the file. Is it corrupted or password-protected? ({exc})")

    # Some PDFs store bullet symbols as unknown glyphs, which come out as
    # "(cid:127)". Turn them back into a normal bullet so we can count them.
    text = re.sub(r"\(cid:\d+\)", "•", text)

    if len(text.strip()) < 50:
        raise FileError(
            "Very little text could be extracted. If your resume is a scanned image, "
            "export it again as a text-based PDF — real ATS systems can't read images either."
        )
    return text
